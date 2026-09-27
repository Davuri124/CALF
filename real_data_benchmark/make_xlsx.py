import pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

OUT_DIR = Path(__file__).parent / "outputs"
res_df = pd.read_csv(OUT_DIR / "results.csv")
res_df = res_df[["Dataset", "Accuracy", "Precision", "Recall", "F1", "ROC-AUC", "N"]]

wb = Workbook()
ws = wb.active
ws.title = "Results"

HEADER_FILL = PatternFill("solid", fgColor="2E5EAA")
AVG_FILL = PatternFill("solid", fgColor="DCE6F1")
FONT_NAME = "Arial"
thin = Side(style="thin", color="B7C4D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

ws["A1"] = "Classification Benchmark Results — Real Public Datasets"
ws["A1"].font = Font(name=FONT_NAME, size=14, bold=True, color="2E5EAA")
ws.merge_cells("A1:G1")
ws["A2"] = "Model: HistGradientBoostingClassifier | 5-fold stratified cross-validation | Metrics in %, ROC-AUC as probability"
ws["A2"].font = Font(name=FONT_NAME, size=9, italic=True, color="666666")
ws.merge_cells("A2:G2")

start_row = 4
headers = list(res_df.columns)
for j, h in enumerate(headers, start=1):
    c = ws.cell(row=start_row, column=j, value=h)
    c.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = border

SMOTE_FILL = PatternFill("solid", fgColor="FFF3CD")
for i, row in res_df.iterrows():
    r = start_row + 1 + i
    is_avg = row["Dataset"] == "Average"
    is_smote = "SMOTE" in str(row["Dataset"])
    for j, h in enumerate(headers, start=1):
        val = row[h]
        c = ws.cell(row=r, column=j, value=val)
        c.font = Font(name=FONT_NAME, bold=is_avg, italic=is_smote)
        c.border = border
        c.alignment = Alignment(horizontal="center" if j > 1 else "left")
        if is_avg:
            c.fill = AVG_FILL
        elif is_smote:
            c.fill = SMOTE_FILL
        if h in ("Accuracy", "Precision", "Recall", "F1"):
            c.number_format = "0.0"
        elif h == "ROC-AUC":
            c.number_format = "0.000"

widths = [34, 12, 12, 10, 10, 11, 8]
for j, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(j)].width = w

ws.freeze_panes = "A5"

# Notes sheet
notes = wb.create_sheet("Notes")
notes["A1"] = "Data Sources"
notes["A1"].font = Font(name=FONT_NAME, size=12, bold=True)
sources = [
    ("Heart Disease", "UCI Cleveland Heart Disease dataset (303 patients, 13 clinical features)"),
    ("Diabetes", "Pima Indians Diabetes dataset (768 patients, 8 diagnostic measurements)"),
    ("Stroke", "Healthcare Stroke Prediction dataset, fedesoriano/Kaggle (5,110 patients)"),
    ("Insurance", "Auto Insurance Fraud Claims dataset, fraud_reported target (1,000 policies)"),
    ("Breast Cancer", "UCI/Wisconsin Diagnostic Breast Cancer dataset (569 patients, via scikit-learn)"),
    ("CKD", "UCI Chronic Kidney Disease dataset (400 patients, 24 clinical attributes)"),
]
notes.append(["Dataset", "Source"])
for a, b in sources:
    notes.append([a, b])
notes["A2"].font = Font(name=FONT_NAME, bold=True)
notes["B2"].font = Font(name=FONT_NAME, bold=True)
for row in notes.iter_rows(min_row=3, max_row=notes.max_row):
    for c in row:
        if c.value is not None:
            c.font = Font(name=FONT_NAME)
notes.column_dimensions["A"].width = 16
notes.column_dimensions["B"].width = 80

notes["A10"] = "Methodology"
notes["A10"].font = Font(name=FONT_NAME, size=12, bold=True)
notes["A11"] = ("Each dataset was cleaned (missing-value handling, categorical typing) and evaluated with a "
                "HistGradientBoostingClassifier under 5-fold stratified cross-validation. Accuracy/Precision/"
                "Recall/F1 are computed at a 0.5 probability threshold on out-of-fold predictions; ROC-AUC uses "
                "out-of-fold predicted probabilities.")
notes["A11"].alignment = Alignment(wrap_text=True)
notes.merge_cells("A11:F14")

notes["A16"] = "Caveat & Stroke Supplementary Row"
notes["A16"].font = Font(name=FONT_NAME, size=12, bold=True)
notes["A17"] = ("Stroke has ~5% positive-class prevalence, so Precision/Recall/F1 are naturally lower than "
                 "Accuracy/ROC-AUC for that dataset alone — this reflects genuine class imbalance in the real "
                 "data, not a modeling error. The highlighted 'Stroke (SMOTE-balanced)' row applies SMOTENC "
                 "oversampling inside each training fold only (test data is never touched) and picks a "
                 "decision threshold that maximizes F1 on the real, imbalanced label distribution. Recall "
                 "rises from 34.9% to 62.0% (the model catches far more true stroke cases), at the cost of "
                 "lower precision — a standard, honest trade-off for medical screening use cases where missing "
                 "a positive case is costlier than a false alarm. This row is a supplementary comparison, not a "
                 "7th dataset, and is excluded from the Average row.")
notes["A17"].alignment = Alignment(wrap_text=True)
notes.merge_cells("A17:F22")

# Baselines sheet
baseline_df = pd.read_csv(OUT_DIR / "baseline_comparison.csv")
bl = wb.create_sheet("Baselines")
bl["A1"] = "Baseline Model Comparison — Logistic Regression vs. Random Forest vs. HistGradientBoosting"
bl["A1"].font = Font(name=FONT_NAME, size=13, bold=True, color="2E5EAA")
bl.merge_cells("A1:G1")
bl_headers = list(baseline_df.columns)
for j, h in enumerate(bl_headers, start=1):
    c = bl.cell(row=3, column=j, value=h)
    c.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center")
    c.border = border
for i, row in baseline_df.iterrows():
    r = 4 + i
    for j, h in enumerate(bl_headers, start=1):
        c = bl.cell(row=r, column=j, value=row[h])
        c.border = border
        c.font = Font(name=FONT_NAME, bold=(h == "Model" and "ours" in str(row[h])))
        c.alignment = Alignment(horizontal="center" if j > 2 else "left")
        if h in ("Accuracy", "Precision", "Recall", "F1"):
            c.number_format = "0.0"
        elif h == "ROC-AUC":
            c.number_format = "0.000"
bl_widths = [16, 26, 12, 12, 10, 10, 11]
for j, w in enumerate(bl_widths, start=1):
    bl.column_dimensions[get_column_letter(j)].width = w
bl.freeze_panes = "A4"

# Charts sheet
charts = wb.create_sheet("Charts")
imgs = ["fig1_metrics_by_dataset.png", "fig4_rocauc_comparison.png", "fig2_roc_curves.png",
        "fig3_confusion_matrices.png", "fig5_stroke_raw_vs_balanced.png", "fig6_pr_curves.png",
        "fig7_error_bars.png", "fig8_baseline_comparison.png", "fig9_heart_feature_importance.png",
        "fig10_ckd_feature_importance.png", "fig11_dataset_characteristics.png"]
row_cursor = 1
for img_name in imgs:
    img = XLImage(str(OUT_DIR / img_name))
    img.width = 640
    img.height = int(640 * img.height / img.width) if img.width else 400
    charts.add_image(img, f"A{row_cursor}")
    row_cursor += 33

wb.save(OUT_DIR / "results.xlsx")
print("Saved results.xlsx")
